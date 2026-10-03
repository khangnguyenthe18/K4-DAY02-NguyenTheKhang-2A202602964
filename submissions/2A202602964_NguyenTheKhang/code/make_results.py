"""make_results.py - Bước 5: tạo results.xlsx, curves/<exp_id>_<mota>.png và các biểu đồ tổng hợp.

Mọi con số đọc từ log chạy thật (runs/*/seed*/summary.json, history.csv, step3.json, final.json) và
từ các file predictions/ (tính bằng eval.compute_metrics của repo gốc), nên xlsx khớp eval.py.
    python make_results.py --work work --labels data/labels --submission ..
"""
from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

import numpy as np
import pandas as pd

import dataset as D
from experiments import (ABLATIONS, AXIS_NAME, BACKBONES, FAMILY, FINAL_SEEDS, T00_SEEDS, Paths, load_decisions,
                         load_summary)


def _plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    return plt


def _j(p: Path):
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def pm(mean, std, d=4):
    return f"{mean:.{d}f} ± {std:.{d}f}"


# --------------------------------------------------------------------------- #
# Đọc kết quả
# --------------------------------------------------------------------------- #
def all_runs(paths: Paths) -> dict[tuple[str, int], dict]:
    out = {}
    for p in Path(paths.out_dir).glob("*/seed*/summary.json"):
        s = json.loads(p.read_text(encoding="utf-8"))
        out[(s["exp_id"], int(s["seed"]))] = s
    return out


def failed_runs(paths: Paths) -> list[str]:
    return sorted(f"{p.parent.parent.name}/{p.parent.name}" for p in Path(paths.out_dir).glob("*/seed*/FAILED.txt"))


def pred_metrics(path: Path) -> dict:
    from eval import compute_metrics, read_pred
    p = read_pred(str(path))
    m = compute_metrics(p.y_true, p.y_pred, p.probs)
    return m


# --------------------------------------------------------------------------- #
# Sheets
# --------------------------------------------------------------------------- #
def sheet_backbones(paths: Paths, runs: dict, decisions: dict) -> pd.DataFrame:
    rows = []
    chosen = decisions.get("backbone", {}).get("backbone")
    for e, name, fam, desc in BACKBONES + [("B08", "vit_small_patch14_dinov2.lvd142m", "Bonus: DINOv2 đóng băng",
                                            "dinov2_probe")]:
        s = runs.get((e, 0))
        if not s:
            continue
        lat = s.get("latency_b1_fp32") or {}
        rows.append({
            "exp_id": e, "backbone": name, "họ": fam, "tag trọng số": s["weight_tag"],
            "#tham số (M)": s["params_M"], "GMAC": s["gmacs"], "độ phân giải (px)": s["img_size"],
            "epoch": s["epochs"], "best epoch": s["best_epoch"], "seed": s["seed"],
            "macro-F1 val": s["val_macro_f1"], "top-1 val": s["val_top1"], "balanced acc val": s["val_bal_acc"],
            "F1 Chinee val": s["val_f1_per_class"][0], "F1 Snake val": s["val_f1_per_class"][7],
            "thời gian train/epoch (s)": s["train_time_per_epoch_s"],
            "độ trễ batch-1 p50 (ms)": lat.get("p50_ms"), "độ trễ batch-1 p95 (ms)": lat.get("p95_ms"),
            "GPU": lat.get("gpu"),
            "ghi chú": ("ĐƯỢC CHỌN cho Bước 2-4. " if name == chosen else "")
                       + ("linear probe, backbone đóng băng, 30 epoch trên đặc trưng" if e == "B08" else
                          "công thức nền T00, 1 seed; độ trễ sơ bộ FP32, 50 lần đo, có warmup + synchronize"),
        })
    return pd.DataFrame(rows)


def sheet_training(paths: Paths, runs: dict, decisions: dict) -> pd.DataFrame:
    t00 = [runs.get(("T00", s)) for s in (0, 1, 2)]
    t00 = [s for s in t00 if s]
    if not t00:
        return pd.DataFrame()
    f1s = [s["val_macro_f1"] for s in t00]
    mean, std = float(np.mean(f1s)), float(np.std(f1s, ddof=1)) if len(f1s) > 1 else float("nan")
    base0 = runs[("T00", 0)]["val_macro_f1"]
    from experiments import verdict
    rows = []

    def rare(s):
        f = s["val_f1_per_class"]
        return {"F1 Chinee val": f[0], "F1 Snake val": f[7], "F1 lớp cỏ thấp nhất val": min(f[:8])}

    for s in t00:
        rows.append({"exp_id": "T00", "backbone": s["backbone"], "trục": "-", "tên trục": "nền",
                     "khác T00 ở điểm nào": "công thức nền (GUIDE 1.4)", "seed": s["seed"],
                     "macro-F1 val": s["val_macro_f1"], "top-1 val": s["val_top1"], "Δ vs T00 seed0": s["val_macro_f1"] - base0,
                     "Δ vs mean T00": s["val_macro_f1"] - mean, "std nhiễu (T00, 3 seed)": std, "kết luận": "mốc", **rare(s),
                     "best epoch": s["best_epoch"], "ghi chú": ""})
    combos = {k: v for k, v in decisions.get("combos", {}).items() if isinstance(v, dict)}
    items = [(e, ax, d, diff) for e, ax, d, diff in ABLATIONS] + \
            [(e, "kết hợp", c["desc"], c["diff"]) for e, c in combos.items()]
    for e, ax, d, diff in items:
        s = runs.get((e, 0))
        if not s:
            continue
        note = ""
        if e in combos:
            parts = combos[e]["from"]
            exp_sum = sum(runs[(p, 0)]["val_macro_f1"] - mean for p in parts if (p, 0) in runs)
            note = (f"gộp {', '.join(parts)} (quy tắc {combos[e]['rule']}); tổng Δ riêng lẻ = {exp_sum:+.4f}, "
                    f"Δ thực = {s['val_macro_f1'] - mean:+.4f} -> "
                    + ("cộng dồn" if s['val_macro_f1'] - mean >= 0.8 * exp_sum else "triệt tiêu một phần"))
        rows.append({"exp_id": e, "backbone": s["backbone"], "trục": ax, "tên trục": AXIS_NAME.get(ax, "kết hợp"),
                     "khác T00 ở điểm nào": ", ".join(f"{k}={v}" for k, v in diff.items()), "seed": s["seed"],
                     "macro-F1 val": s["val_macro_f1"], "top-1 val": s["val_top1"],
                     "Δ vs T00 seed0": s["val_macro_f1"] - base0, "Δ vs mean T00": s["val_macro_f1"] - mean,
                     "std nhiễu (T00, 3 seed)": std, "kết luận": verdict(s["val_macro_f1"] - mean, std), **rare(s),
                     "best epoch": s["best_epoch"], "ghi chú": note})
    for s in [runs.get(("F01", k)) for k in FINAL_SEEDS]:
        if s:
            rows.append({"exp_id": "F01", "backbone": s["backbone"], "trục": "chung kết", "tên trục": "chung kết",
                         "khác T00 ở điểm nào": decisions.get("recipe", {}).get("diff_str", ""), "seed": s["seed"],
                         "macro-F1 val": s["val_macro_f1"], "top-1 val": s["val_top1"],
                         "Δ vs T00 seed0": s["val_macro_f1"] - base0, "Δ vs mean T00": s["val_macro_f1"] - mean,
                         "std nhiễu (T00, 3 seed)": std, "kết luận": "", **rare(s), "best epoch": s["best_epoch"],
                         "ghi chú": "1 view, chưa TTA/TS (macro-F1 của checkpoint tốt nhất)"})
    return pd.DataFrame(rows)


def sheet_inference(step3: dict | None) -> pd.DataFrame:
    if not step3:
        return pd.DataFrame()
    cols = {"exp_id": "exp_id", "method": "phương pháp", "model": "mô hình/checkpoint", "K": "K (view/model)",
            "val_macro_f1": "macro-F1 val", "val_top1": "top-1 val", "val_ece": "ECE val (15 bin)",
            "val_f1_chinee": "F1 Chinee val", "val_f1_snake": "F1 Snake val",
            "p50_ms": "p50 b1 (ms)", "p95_ms": "p95 b1 (ms)", "p99_ms": "p99 b1 (ms)",
            "throughput_b32": "thông lượng b32 (ảnh/s)", "rel_cost_vs_I00": "chi phí tương đối vs I00 (p50)",
            "spec_name": "spec", "note": "ghi chú"}
    df = pd.DataFrame(step3["rows"])
    for c in cols:
        if c not in df:
            df[c] = None
    df = df[list(cols)].rename(columns=cols)
    return df


def sheet_final(paths: Paths, decisions: dict, runs: dict) -> tuple[pd.DataFrame, dict]:
    pdir = Path(paths.pred_dir)
    rows, agg = [], {}
    desc = {"F01": decisions.get("recipe", {}).get("summary", "F01"), "T00": "T00 công thức nền + I00 (1 view), không TS"}
    for tag in ("F01", "T00"):
        vals = []
        for k in (FINAL_SEEDS if tag == "F01" else T00_SEEDS):
            tp, vp = pdir / f"{tag}_seed{k}_test.csv", pdir / f"{tag}_seed{k}_val.csv"
            if not tp.exists():
                continue
            mt, mv = pred_metrics(tp), pred_metrics(vp) if vp.exists() else None
            up = pdir / f"F01uncal_seed{k}_test.csv"
            ece_u = pred_metrics(up)["ece"] if (tag == "F01" and up.exists()) else None
            r = {"exp_id": tag, "cấu hình": desc[tag], "seed": k,
                 "macro-F1 val": mv["macro_f1"] if mv else None, "macro-F1 test": mt["macro_f1"],
                 "top-1 test": mt["top1"], "balanced acc test": mt["balanced_acc"], "ECE test": mt["ece"],
                 "ECE test trước TS": ece_u, "recall Chinee test": mt["recall"][0], "recall Snake test": mt["recall"][7],
                 "F1 Chinee test": mt["f1"][0], "F1 Snake test": mt["f1"][7]}
            rows.append(r)
            vals.append(r)
        if vals:
            a = {}
            for c in ("macro-F1 val", "macro-F1 test", "top-1 test", "balanced acc test", "ECE test", "ECE test trước TS",
                      "recall Chinee test", "recall Snake test", "F1 Chinee test", "F1 Snake test"):
                v = [r[c] for r in vals if r[c] is not None]
                if v:
                    a[c] = (float(np.mean(v)), float(np.std(v, ddof=1)) if len(v) > 1 else float("nan"))
            agg[tag] = a
            rows.append({"exp_id": f"{tag} (mean ± std, {len(vals)} seed)", "cấu hình": desc[tag], "seed": "tổng hợp",
                         **{c: pm(*a[c]) for c in a}})
    return pd.DataFrame(rows), agg


def sheet_perclass(paths: Paths) -> pd.DataFrame:
    ev = Path(paths.work_dir) / "eval_out"
    frames = []
    for tag, label in (("F01", "chung kết F01 (sau TS)"), ("T00", "mốc T00 + I00")):
        p = ev / f"{tag}_per_class.csv"
        if p.exists():
            d = pd.read_csv(p)
            frames.append(pd.DataFrame({
                "cấu hình": label, "lớp": d["class"], "số ảnh test": d["support"],
                "precision": [pm(a, b, 3) for a, b in zip(d.precision_mean, d.precision_std)],
                "recall": [pm(a, b, 3) for a, b in zip(d.recall_mean, d.recall_std)],
                "F1": [pm(a, b, 3) for a, b in zip(d.f1_mean, d.f1_std)],
                "F1 mean": d.f1_mean, "recall mean": d.recall_mean}))
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def sheet_latency(step3: dict | None, final: dict | None, runs: dict) -> pd.DataFrame:
    rows = []
    for (e, s), r in sorted(runs.items()):
        if e.startswith("B") and s == 0 and r.get("latency_b1_fp32"):
            rows.append({**r["latency_b1_fp32"], "config": f"{e} {r['backbone']} (sơ bộ, Bước 1)"})
    if step3:
        rows += step3["latency"]
    if final:
        rows += [final["final_latency"]]
        if final["realtime_latency"] is not final["final_latency"]:
            rows.append(final["realtime_latency"])
    cols = ["config", "gpu", "dtype", "batch", "img_size", "fused_bn", "channels_last", "k_views", "tta_mode",
            "p50_ms", "p95_ms", "p99_ms", "mean_ms", "images_per_s", "n_iters", "warmup", "torch", "preprocessing_included"]
    df = pd.DataFrame(rows)
    for c in cols:
        if c not in df:
            df[c] = None
    df = df[cols].drop_duplicates()
    return df.rename(columns={"config": "cấu hình", "gpu": "GPU", "dtype": "dtype", "batch": "batch",
                              "img_size": "độ phân giải (px)", "fused_bn": "gộp BN", "p50_ms": "p50 (ms)",
                              "p95_ms": "p95 (ms)", "p99_ms": "p99 (ms)", "mean_ms": "mean (ms)",
                              "images_per_s": "ảnh/s", "n_iters": "số lần đo", "warmup": "warmup",
                              "preprocessing_included": "tính tiền xử lý"})


def sheet_summary(bb: pd.DataFrame, tr: pd.DataFrame, inf: pd.DataFrame, agg: dict, final: dict | None) -> pd.DataFrame:
    cand = []
    for _, r in bb.iterrows():
        cand.append({"exp_id": r["exp_id"], "loại": "backbone", "mô tả": r["backbone"], "macro-F1 val": r["macro-F1 val"],
                     "top-1 val": r["top-1 val"], "p50 b1 (ms)": r["độ trễ batch-1 p50 (ms)"], "GMAC": r["GMAC"]})
    if len(tr):
        for _, r in tr[tr["seed"] == 0].iterrows():
            if r["exp_id"] in ("T00", "F01"):
                continue
            cand.append({"exp_id": r["exp_id"], "loại": "công thức", "mô tả": r["khác T00 ở điểm nào"],
                         "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"]})
    if len(inf):
        for _, r in inf.iterrows():
            if pd.notna(r["macro-F1 val"]):
                cand.append({"exp_id": r["exp_id"], "loại": "suy luận", "mô tả": r["phương pháp"],
                             "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"],
                             "p50 b1 (ms)": r["p50 b1 (ms)"], "chi phí vs I00": r["chi phí tương đối vs I00 (p50)"]})
    top = pd.DataFrame(cand).sort_values("macro-F1 val", ascending=False).head(10)
    top.insert(0, "hạng", range(1, len(top) + 1))
    rows = [{"hạng": "TOP 10 THEO MACRO-F1 VAL (1 seed, trừ khi ghi khác)"}] + top.to_dict("records")
    if agg:
        rows.append({})
        rows.append({"hạng": "CHUNG KẾT vs MỐC (TEST, mean ± std qua 3 seed, tính từ predictions/ bằng eval.py)"})
        for tag, name in (("T00", "Mốc: T00 + I00"), ("F01", "Chung kết F01")):
            if tag in agg:
                a = agg[tag]
                rows.append({"hạng": name, "exp_id": tag, "macro-F1 val": pm(*a["macro-F1 val"]) if "macro-F1 val" in a else "",
                             "mô tả": f"macro-F1 test {pm(*a['macro-F1 test'])} | top-1 test {pm(*a['top-1 test'])} | "
                                      f"ECE test {pm(*a['ECE test'])} | recall Chinee {pm(*a['recall Chinee test'], 3)} | "
                                      f"recall Snake {pm(*a['recall Snake test'], 3)}"})
        if "F01" in agg and "T00" in agg:
            d = agg["F01"]["macro-F1 test"][0] - agg["T00"]["macro-F1 test"][0]
            s = max(agg["F01"]["macro-F1 test"][1], agg["T00"]["macro-F1 test"][1])
            rows.append({"hạng": "Cải thiện", "mô tả": f"Δ macro-F1 test = {d:+.4f}; std lớn hơn = {s:.4f}; "
                                                      f"{'Δ > std' if d > s else 'Δ ≤ std (không phân biệt được)'}"})
        if final:
            rl = final["realtime_latency"]
            rows.append({"hạng": "Thời gian thực", "mô tả": f"{rl['config']}: p95 batch-1 = {rl['p95_ms']:.2f} ms "
                                                           f"trên {rl['gpu']} ({rl['dtype']})"})
    return pd.DataFrame(rows)


def sheet_sanity(step0: dict | None) -> pd.DataFrame:
    if not step0:
        return pd.DataFrame()
    s, e = step0["sanity"], step0["eda"]
    rows = [{"kiểm tra": "Giao train∩val / train∩test / val∩test", "kết quả": str(e["split_check"]["overlap"]), "đạt": True},
            {"kiểm tra": "Hợp ba tập = 17.509", "kết quả": e["split_check"]["union"],
             "đạt": e["split_check"]["union"] == 17509},
            {"kiểm tra": "Số ảnh train/val/test", "kết quả": str(e["split_check"]["n"]), "đạt": True},
            {"kiểm tra": "File thiếu trên đĩa", "kết quả": e["split_check"]["missing_files"],
             "đạt": e["split_check"]["missing_files"] == 0}]
    for k, v in s["initial_loss"].items():
        rows.append({"kiểm tra": f"Loss ban đầu {k} (kỳ vọng ≈ ln 9 = 2.197)", "kết quả": round(v, 4),
                     "đạt": abs(v - 2.197) < 0.35})
    o = s["overfit"]
    rows.append({"kiểm tra": f"Overfit {o['n_images']} ảnh / {o['steps']} bước", "kết quả":
                 f"loss {o['loss_first']:.3f} -> {o['loss_last']:.4f}, acc {o['train_acc_eval_mode']:.2f}", "đạt": o["passed"]})
    rows.append({"kiểm tra": "BN ở eval khi đóng băng backbone", "kết quả": str(s["frozen_bn_eval"]),
                 "đạt": s["frozen_bn_eval"]["n_bn_training"] == 0})
    rows.append({"kiểm tra": "|Focal(γ=0) − CE|", "kết quả": s["focal_gamma0_minus_ce"], "đạt": s["focal_gamma0_minus_ce"] < 1e-6})
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------- #
# Ghi xlsx có định dạng
# --------------------------------------------------------------------------- #
def write_xlsx(sheets: dict[str, pd.DataFrame], path: Path, highlight: dict[str, str]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            (df if len(df) else pd.DataFrame({"ghi chú": ["chưa có dữ liệu"]})).to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "B2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2F5597")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for i, col in enumerate(df.columns if len(df) else ["ghi chú"], 1):
                width = max([len(str(col))] + [len(str(v)) for v in (df[col].head(50) if len(df) else [])])
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    for c in row:
                        if isinstance(c.value, float):
                            c.number_format = "0.00" if "(ms)" in str(col) or "(s)" in str(col) or "ảnh/s" in str(col) \
                                else "0.0000"
            hcol = highlight.get(name)
            if hcol and len(df) and hcol in df and pd.to_numeric(df[hcol], errors="coerce").notna().any():
                best = pd.to_numeric(df[hcol], errors="coerce").idxmax() + 2
                for c in ws[best]:
                    c.fill = PatternFill("solid", fgColor="FFF2CC")
                    c.font = Font(bold=True)


# --------------------------------------------------------------------------- #
# Biểu đồ
# --------------------------------------------------------------------------- #
def export_curves(paths: Paths, runs: dict, curves_dir: Path) -> list[str]:
    """Mỗi exp_id một ảnh curves/<exp_id>_<desc>.png; exp nhiều seed vẽ chồng các seed."""
    plt = _plt()
    curves_dir.mkdir(parents=True, exist_ok=True)
    by_exp: dict[str, list[dict]] = {}
    for (e, s), r in runs.items():
        by_exp.setdefault(e, []).append(r)
    written = []
    for e, rs in sorted(by_exp.items()):
        desc = rs[0].get("desc") or rs[0]["backbone"]
        dst = curves_dir / f"{e}_{desc}.png"
        if len(rs) == 1:
            shutil.copy(Path(paths.out_dir) / e / f"seed{rs[0]['seed']}" / "curve.png", dst)
        else:
            fig, ax = plt.subplots(1, 3, figsize=(16, 4.2))
            for r in sorted(rs, key=lambda r: r["seed"]):
                h = pd.read_csv(Path(paths.out_dir) / e / f"seed{r['seed']}" / "history.csv")
                lr = np.load(Path(paths.out_dir) / e / f"seed{r['seed']}" / "lr_steps.npy")
                ax[0].plot(h.epoch, h.train_loss, "o-", label=f"train seed{r['seed']}")
                ax[0].plot(h.epoch, h.val_loss, "s--", label=f"val seed{r['seed']}")
                ax[1].plot(h.epoch, h.val_macro_f1, "o-", label=f"val macro-F1 seed{r['seed']}")
                ax[1].plot(h.epoch, h.val_top1, "s:", alpha=.6, label=f"val top-1 seed{r['seed']}")
                ax[2].plot(lr, label=f"seed{r['seed']}")
            ax[0].set(xlabel="epoch", ylabel="loss", title="Loss (train / val CE)")
            ax[1].set(xlabel="epoch", ylabel="metric", title="Val metrics")
            ax[2].set(xlabel="step", ylabel="LR (backbone)", title="LR schedule")
            for a in ax:
                a.grid(alpha=.3)
                a.legend(fontsize=7)
            fig.suptitle(f"{e} | {rs[0]['backbone']} | {len(rs)} seed | {desc}")
            fig.tight_layout()
            fig.savefig(dst, dpi=130)
            plt.close(fig)
        written.append(dst.name)
    return written


def summary_figures(paths: Paths, bb: pd.DataFrame, tr: pd.DataFrame, step3: dict | None, fig_dir: Path) -> None:
    plt = _plt()
    if len(bb):
        fig, ax = plt.subplots(figsize=(7, 4.5))
        for _, r in bb.iterrows():
            ax.scatter(r["độ trễ batch-1 p50 (ms)"], r["macro-F1 val"], s=20 + 8 * r["#tham số (M)"], alpha=.6)
            ax.annotate(f"{r['exp_id']} {r['backbone'].split('.')[0]}\n{r['GMAC']:.2f} GMAC",
                        (r["độ trễ batch-1 p50 (ms)"], r["macro-F1 val"]), fontsize=7, xytext=(4, 4),
                        textcoords="offset points")
        ax.set(xlabel="độ trễ batch-1 p50 (ms, FP32)", ylabel="macro-F1 val",
               title="Bước 1: chất lượng vs độ trễ (kích thước điểm ~ #tham số)")
        ax.grid(alpha=.3)
        fig.tight_layout()
        fig.savefig(fig_dir / "fig_backbones_f1_latency.png", dpi=130)
        plt.close(fig)
    if len(tr):
        t = tr[(tr.seed == 0) & (~tr.exp_id.isin(["T00", "F01"]))]
        std = tr["std nhiễu (T00, 3 seed)"].iloc[0]
        fig, ax = plt.subplots(figsize=(10, 4.5))
        colors = ["tab:green" if d > std else ("tab:red" if d < -std else "tab:gray") for d in t["Δ vs mean T00"]]
        ax.bar(t.exp_id + "\n" + t["khác T00 ở điểm nào"].str.slice(0, 18), t["Δ vs mean T00"], color=colors)
        ax.axhspan(-std, std, color="gray", alpha=.2, label=f"±1 std nhiễu T00 ({std:.4f})")
        ax.axhline(0, c="k", lw=.8)
        ax.set(ylabel="Δ macro-F1 val vs mean T00", title="Bước 2: ablation một yếu tố (xanh: > std, đỏ: < −std)")
        ax.tick_params(axis="x", labelsize=6)
        ax.legend()
        ax.grid(alpha=.3, axis="y")
        fig.tight_layout()
        fig.savefig(fig_dir / "fig_ablation_delta.png", dpi=130)
        plt.close(fig)
    if step3:
        rows = [r for r in step3["rows"] if r.get("p50_ms") and r.get("val_macro_f1") is not None]
        fig, ax = plt.subplots(figsize=(8, 5))
        for r in rows:
            ax.scatter(r["p50_ms"], r["val_macro_f1"], s=30)
            ax.annotate(r["exp_id"], (r["p50_ms"], r["val_macro_f1"]), fontsize=7, xytext=(3, 3), textcoords="offset points")
        ax.axvline(100, ls="--", c="r", lw=.8, label="ngân sách 100 ms")
        ax.set(xscale="log", xlabel="độ trễ batch-1 p50 (ms, log)", ylabel="macro-F1 val",
               title="Bước 3: đánh đổi độ chính xác – độ trễ")
        ax.legend()
        ax.grid(alpha=.3)
        fig.tight_layout()
        fig.savefig(fig_dir / "fig_inference_tradeoff.png", dpi=130)
        plt.close(fig)


def final_figures(paths: Paths, fig_dir: Path) -> None:
    from eval import ECE_BINS, read_pred
    plt = _plt()
    ev = Path(paths.work_dir) / "eval_out"
    cm_p = ev / "F01_confusion_sum.csv"
    if cm_p.exists():
        cm = pd.read_csv(cm_p, index_col=0).to_numpy()
        cmn = cm / cm.sum(1, keepdims=True)
        fig, ax = plt.subplots(figsize=(8, 7))
        ax.imshow(cmn, cmap="Blues", vmin=0, vmax=1)
        for i in range(9):
            for j in range(9):
                ax.text(j, i, f"{cm[i, j]}\n{cmn[i, j]:.1%}", ha="center", va="center", fontsize=6.5,
                        color="white" if cmn[i, j] > .5 else "black")
        names = [n[:12] for n in D.CLASS_NAMES]
        ax.set_xticks(range(9), names, rotation=40, ha="right", fontsize=8)
        ax.set_yticks(range(9), names, fontsize=8)
        ax.set(xlabel="dự đoán", ylabel="nhãn thật", title="Ma trận nhầm lẫn TEST, F01 (cộng 3 seed; % theo hàng)")
        fig.tight_layout()
        fig.savefig(fig_dir / "fig_confusion_test.png", dpi=130)
        plt.close(fig)
    pdir = Path(paths.pred_dir)
    s0 = FINAL_SEEDS[0]
    if (pdir / f"F01_seed{s0}_test.csv").exists() and (pdir / f"F01uncal_seed{s0}_test.csv").exists():
        fig, ax = plt.subplots(figsize=(5.5, 5))
        for tag, lab in (("F01uncal", "trước TS"), ("F01", "sau TS")):
            p = read_pred(str(pdir / f"{tag}_seed{s0}_test.csv"))
            conf, corr = p.probs.max(1), (p.y_pred == p.y_true)
            idx = np.clip(np.ceil(conf * ECE_BINS).astype(int) - 1, 0, ECE_BINS - 1)
            xs, ys = [], []
            for b in range(ECE_BINS):
                m = idx == b
                if m.sum() >= 5:
                    xs.append(conf[m].mean())
                    ys.append(corr[m].mean())
            ax.plot(xs, ys, "o-", label=lab)
        ax.plot([0, 1], [0, 1], "k:", label="hiệu chuẩn hoàn hảo")
        ax.set(xlabel="độ tin cậy (max softmax)", ylabel="accuracy trong bin", xlim=(0.3, 1.01), ylim=(0.3, 1.01),
               title=f"Reliability diagram TEST (F01 seed {s0}, 15 bin)")
        ax.legend()
        ax.grid(alpha=.3)
        fig.tight_layout()
        fig.savefig(fig_dir / "fig_reliability_test.png", dpi=130)
        plt.close(fig)


# --------------------------------------------------------------------------- #
def build_all(paths: Paths, submission_dir: Path) -> dict:
    runs = all_runs(paths)
    decisions = load_decisions(paths)
    step0 = _j(Path(paths.work_dir) / "step0.json")
    step3 = _j(Path(paths.work_dir) / "step3.json")
    final = _j(Path(paths.work_dir) / "final.json")
    bb = sheet_backbones(paths, runs, decisions)
    tr = sheet_training(paths, runs, decisions)
    inf = sheet_inference(step3)
    fin, agg = sheet_final(paths, decisions, runs)
    pc = sheet_perclass(paths)
    lat = sheet_latency(step3, final, runs)
    summ = sheet_summary(bb, tr, inf, agg, final)
    san = sheet_sanity(step0)
    sheets = {"Summary": summ, "Backbones": bb, "Training": tr, "Inference": inf, "Final": fin,
              "PerClass": pc, "Latency": lat, "Sanity": san}
    submission_dir.mkdir(parents=True, exist_ok=True)
    write_xlsx(sheets, submission_dir / "results.xlsx",
               {"Backbones": "macro-F1 val", "Training": "macro-F1 val", "Inference": "macro-F1 val"})
    fig_dir = submission_dir / "figures"
    fig_dir.mkdir(exist_ok=True)
    for f in Path(paths.work_dir, "figures").glob("*.png"):
        shutil.copy(f, fig_dir / f.name)
    summary_figures(paths, bb, tr, step3, fig_dir)
    final_figures(paths, fig_dir)
    curves = export_curves(paths, runs, submission_dir / "curves")
    # predictions/ (chỉ chung kết + mốc) và đầu ra eval.py
    if Path(paths.pred_dir).exists():
        shutil.copytree(paths.pred_dir, submission_dir / "predictions", dirs_exist_ok=True)
    if (Path(paths.work_dir) / "eval_out").exists():
        shutil.copytree(Path(paths.work_dir) / "eval_out", submission_dir / "eval_out", dirs_exist_ok=True)
    # log nhẹ để truy vết exp_id -> log (không chép checkpoint)
    logs = submission_dir / "logs"
    for p in Path(paths.out_dir).glob("*/seed*"):
        dst = logs / p.parent.name / p.name
        dst.mkdir(parents=True, exist_ok=True)
        for f in ("config.json", "summary.json", "history.csv", "FAILED.txt"):
            if (p / f).exists():
                shutil.copy(p / f, dst / f)
    for f in ("decisions.json", "step0.json", "step3.json", "final.json"):
        if (Path(paths.work_dir) / f).exists():
            shutil.copy(Path(paths.work_dir) / f, logs / f)
    if (Path(paths.out_dir) / "test_access.log").exists():
        shutil.copy(Path(paths.out_dir) / "test_access.log", logs / "test_access.log")
    print(f"results.xlsx: {len(sheets)} sheet; curves: {len(curves)} ảnh; thất bại: {failed_runs(paths)}")
    return {"sheets": {k: len(v) for k, v in sheets.items()}, "curves": curves, "agg": agg,
            "failed": failed_runs(paths)}


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--work", default="work")
    ap.add_argument("--images", default="data/images")
    ap.add_argument("--labels", default="data/labels")
    ap.add_argument("--submission", default="..")
    a = ap.parse_args()
    build_all(Paths(images_dir=a.images, labels_dir=a.labels, work_dir=a.work), Path(a.submission))
